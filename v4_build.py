#!/usr/bin/env python3
"""
V4. Assemble the TTS manuscript: parse the text, build the tables from the
result files, number the references by first appearance, render the .docx and
convert it to PDF for reading.
"""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from v3_docx import (OUT, ROOT, SERIF, _cols, _margins, _para, _runs,
                     equation, figure, table, wide_block)

Q = Path("/home/claude/q1/results")
FX = Path("/home/claude/fix/results")
SF = Path("/home/claude/scutfix/results")
SR = Path("/home/claude/scutrun/results")
TR = ROOT / "results"

AFFIL = (
    "Q. M. Rizvi is with the Department of CSE-CSDS and is Head of the "
    "Research and Development Cell, Brainware University, Barasat, Kolkata, "
    "West Bengal 700125, India (e-mail: sirqaim@gmail.com). "
    "T. Fatima is a Research Scholar with the Department of Biotechnology and "
    "Bioengineering, Khwaja Moinuddin Chishti Language University, Lucknow, "
    "India (e-mail: tehzeebfatima1272@gmail.com). "
    "S. Talat is a Research Scholar with the Department of CSE-CSDS, "
    "Brainware University, Kolkata, India."
)

ETHICS = (
    "An ethics determination for this study is recorded under "
    "[REFERENCE NUMBER], issued [DATE] by the institutional review body of "
    "Brainware University. Code and the full analysis pipeline are available "
    "at [REPOSITORY URL], archived at [DOI]."
)

REFS = {
    "buolamwini2018": "J. Buolamwini and T. Gebru, “Gender shades: intersectional accuracy disparities in commercial gender classification,” in Proc. ACM Conf. Fairness, Accountability, Transparency (FAccT), PMLR 81, 2018, pp. 77–91.",
    "raji2019": "I. D. Raji and J. Buolamwini, “Actionable auditing: investigating the impact of publicly naming biased performance results of commercial AI products,” in Proc. AAAI/ACM Conf. AI, Ethics, and Society (AIES), 2019, pp. 429–435.",
    "grother2019": "P. Grother, M. Ngan, and K. Hanaoka, “Face recognition vendor test (FRVT) part 3: demographic effects,” Nat. Inst. Standards Technol., Gaithersburg, MD, USA, NISTIR 8280, 2019.",
    "mitchell2019": "M. Mitchell et al., “Model cards for model reporting,” in Proc. ACM Conf. Fairness, Accountability, Transparency (FAccT), 2019, pp. 220–229.",
    "gebru2021": "T. Gebru et al., “Datasheets for datasets,” Commun. ACM, vol. 64, no. 12, pp. 86–92, 2021.",
    "euaia2024": "European Parliament and Council of the European Union, “Regulation (EU) 2024/1689 laying down harmonised rules on artificial intelligence (Artificial Intelligence Act),” Off. J. Eur. Union, Jun. 2024.",
    "nist2023": "National Institute of Standards and Technology, “Artificial intelligence risk management framework (AI RMF 1.0),” NIST AI 100-1, Jan. 2023.",
    "spearman1904": "C. Spearman, “The proof and measurement of association between two things,” Amer. J. Psychol., vol. 15, no. 1, pp. 72–101, 1904.",
    "thorndike1949": "R. L. Thorndike, Personnel Selection: Test and Measurement Techniques. New York, NY, USA: Wiley, 1949.",
    "hunter2015": "F. L. Schmidt and J. E. Hunter, Methods of Meta-Analysis: Correcting Error and Bias in Research Findings, 3rd ed. Thousand Oaks, CA, USA: Sage, 2015.",
    "kotwal2025": "K. Kotwal and S. Marcel, “Review of demographic fairness in face recognition,” IEEE Trans. Biometrics, Behav., Identity Sci., 2025.",
    "karkkainen2021": "K. Kärkkäinen and J. Joo, “FairFace: face attribute dataset for balanced race, gender, and age,” in Proc. IEEE Winter Conf. Appl. Comput. Vis. (WACV), 2021, pp. 1548–1558.",
    "jacobs2021": "A. Z. Jacobs and H. Wallach, “Measurement and fairness,” in Proc. ACM Conf. Fairness, Accountability, Transparency (FAccT), 2021, pp. 375–385.",
    "fogliato2020": "R. Fogliato, A. Chouldechova, and M. G’Sell, “Fairness evaluation in presence of biased noisy labels,” in Proc. Int. Conf. Artif. Intell. Statist. (AISTATS), PMLR 108, 2020, pp. 2325–2336.",
    "mitchell2021": "S. Mitchell, E. Potash, S. Barocas, A. D’Amour, and K. Lum, “Algorithmic fairness: choices, assumptions, and definitions,” Annu. Rev. Statist. Appl., vol. 8, pp. 141–163, 2021.",
    "chouldechova2017": "A. Chouldechova, “Fair prediction with disparate impact: a study of bias in recidivism prediction instruments,” Big Data, vol. 5, no. 2, pp. 153–163, 2017.",
    "laurentini2014": "A. Laurentini and A. Bottino, “Computer analysis of face beauty: a survey,” Comput. Vis. Image Underst., vol. 125, pp. 184–199, 2014.",
    "liu2016": "S. Liu et al., “Advances in computational facial attractiveness methods,” Multimedia Tools Appl., vol. 75, no. 23, pp. 16633–16663, 2016.",
    "eisenthal2006": "Y. Eisenthal, G. Dror, and E. Ruppin, “Facial attractiveness: beauty and the machine,” Neural Comput., vol. 18, no. 1, pp. 119–142, 2006.",
    "kagian2006": "A. Kagian et al., “A humanlike predictor of facial attractiveness,” in Adv. Neural Inf. Process. Syst., vol. 19, 2006, pp. 649–656.",
    "schmid2008": "K. Schmid, D. Marx, and A. Samal, “Computation of a face attractiveness index based on neoclassical canons, symmetry, and golden ratio,” Pattern Recognit., vol. 41, no. 8, pp. 2710–2717, 2008.",
    "zhang2011": "D. Zhang, Q. Zhao, and F. Chen, “Quantitative analysis of human facial beauty using geometric features,” Pattern Recognit., vol. 44, no. 4, pp. 940–950, 2011.",
    "rizvi2013": "Q. M. Rizvi, A. Karawia, and S. Kumar, “Female facial beauty analysis for assessment of facial attractiveness,” in Proc. IEEE Int. Conf. Inf. Manage. Knowl. Eng., 2013.",
    "gray2010": "D. Gray et al., “Predicting facial beauty without landmarks,” in Proc. Eur. Conf. Comput. Vis. (ECCV), 2010, pp. 434–447.",
    "fan2018": "Y.-Y. Fan et al., “Label distribution-based facial attractiveness computation by deep residual learning,” IEEE Trans. Multimedia, vol. 20, no. 8, pp. 2196–2208, 2018.",
    "lin2019": "L. Lin et al., “Attribute-aware convolutional neural networks for facial beauty prediction,” in Proc. Int. Joint Conf. Artif. Intell. (IJCAI), 2019, pp. 847–853.",
    "xu2020": "L. Xu and J. Xiang, “ComboLoss for facial attractiveness analysis with squeeze-and-excitation networks,” arXiv:2010.10721, 2020.",
    "vmbeauty2025": "“VM-BeautyNet: a synergistic ensemble of vision transformer and Mamba for facial beauty prediction,” arXiv:2510.16220, 2025.",
    "liang2018": "L. Liang et al., “SCUT-FBP5500: a diverse benchmark dataset for multi-paradigm facial beauty prediction,” in Proc. Int. Conf. Pattern Recognit. (ICPR), 2018, pp. 1598–1603.",
    "geng2013": "X. Geng, C. Yin, and Z.-H. Zhou, “Facial age estimation by learning from label distributions,” IEEE Trans. Pattern Anal. Mach. Intell., vol. 35, no. 10, pp. 2401–2412, 2013.",
    "lebedeva2022": "I. Lebedeva, Y. Guo, and F. Ying, “MEBeauty: a multi-ethnic facial beauty dataset in-the-wild,” Neural Comput. Appl., vol. 34, pp. 14169–14183, 2022.",
    "mebench2026": "I. Lebedeva et al., “MEBeauty-Benchmark.” [Online]. Available: https://github.com/dr-irina-lebedeva/MEBeauty-Benchmark",
    "hamel2025": "C. Hamel and M. Busch, “Analysis of bias in deep learning facial beauty regressors,” arXiv:2509.24138, 2025.",
    "holland2008": "E. Holland, “Marquardt’s Phi Mask: pitfalls of relying on fashion models and the golden ratio to describe a beautiful face,” Aesthetic Plast. Surg., vol. 32, no. 2, pp. 200–208, 2008.",
    "langlois1990": "J. H. Langlois and L. A. Roggman, “Attractive faces are only average,” Psychol. Sci., vol. 1, no. 2, pp. 115–121, 1990.",
    "rhodes1999": "G. Rhodes, A. Sumich, and G. Byatt, “Are average facial configurations attractive only because of their symmetry?” Psychol. Sci., vol. 10, no. 1, pp. 52–58, 1999.",
    "rhodes2006": "G. Rhodes, “The evolution of facial attractiveness,” Annu. Rev. Psychol., vol. 57, pp. 199–226, 2006.",
    "langlois2000": "J. H. Langlois et al., “Maxims or myths of beauty? A meta-analytic and theoretical review,” Psychol. Bull., vol. 126, no. 3, pp. 390–423, 2000.",
    "farkas1985": "L. G. Farkas et al., “Vertical and horizontal proportions of the face in young adult North American Caucasians: revision of neoclassical canons,” Plast. Reconstr. Surg., vol. 75, no. 3, pp. 328–338, 1985.",
    "farkas1994": "L. G. Farkas, Anthropometry of the Head and Face, 2nd ed. New York, NY, USA: Raven Press, 1994.",
    "kazemi2014": "V. Kazemi and J. Sullivan, “One millisecond face alignment with an ensemble of regression trees,” in Proc. IEEE Conf. Comput. Vis. Pattern Recognit. (CVPR), 2014, pp. 1867–1874.",
    "kartynnik2019": "Y. Kartynnik et al., “Real-time facial surface geometry from monocular video on mobile GPUs,” arXiv:1907.06724, 2019.",
    "vaswani2017": "A. Vaswani et al., “Attention is all you need,” in Adv. Neural Inf. Process. Syst., vol. 30, 2017, pp. 5998–6008.",
    "benjamini1995": "Y. Benjamini and Y. Hochberg, “Controlling the false discovery rate: a practical and powerful approach to multiple testing,” J. Roy. Statist. Soc. B, vol. 57, no. 1, pp. 289–300, 1995.",
    "holm1979": "S. Holm, “A simple sequentially rejective multiple test procedure,” Scand. J. Statist., vol. 6, no. 2, pp. 65–70, 1979.",
    "kruskal1952": "W. H. Kruskal and W. A. Wallis, “Use of ranks in one-criterion variance analysis,” J. Amer. Statist. Assoc., vol. 47, no. 260, pp. 583–621, 1952.",
    "cliff1993": "N. Cliff, “Dominance statistics: ordinal analyses to answer ordinal questions,” Psychol. Bull., vol. 114, no. 3, pp. 494–509, 1993.",
}

BIOS = [
    ("Qaim Mehdi Rizvi",
     "received his doctoral degree in computer science and engineering. He is "
     "with the Department of CSE-CSDS, Brainware University, Kolkata, India, "
     "where he is Head of the Research and Development Cell. His research "
     "interests include computer vision, interpretable machine learning, and "
     "the evaluation of automated decision systems."),
    ("Tehzeeb Fatima",
     "is a Research Scholar with the Department of Biotechnology and "
     "Bioengineering, Khwaja Moinuddin Chishti Language University, Lucknow, "
     "India. Her research interests include computational biology and the "
     "statistical analysis of biometric measurements."),
    ("Sana Talat",
     "is a Research Scholar with the Department of CSE-CSDS, Brainware "
     "University, Kolkata, India. Her research interests include machine "
     "learning, facial image analysis, and algorithmic fairness."),
]

MEB_LBL = {"asian": "Asian", "black": "Black", "caucasian": "Caucasian",
           "hispanic": "Hispanic", "indian": "Indian",
           "mideastern": "Mid-Eastern"}
SCUT_LBL = {"asian_female": "Asian F", "asian_male": "Asian M",
            "caucasian_female": "Caucasian F", "caucasian_male": "Caucasian M"}


# ------------------------------------------------------------------ tables
def tbl_validation(doc):
    v = pd.read_csv(TR / "v1_estimator_recovery.csv")
    rows = [["Condition", "Observed", "Corrected", "True \u03c1"]]
    for _, r in v.iterrows():
        rows.append([f"{r['check'].capitalize()}, {r['parameter']}",
                     f"{r['observed']:.4f}", f"{r['corrected']:.4f}",
                     f"{r['true_rho']:.2f}"])
    table(doc, rows, [1.62, 0.62, 0.64, 0.52])


def tbl_floor(doc):
    f = pd.read_csv(TR / "v1_noise_floor.csv")
    rows = [["Evaluation design", "G", "ρ = .60", "ρ = .70",
             "ρ = .80"]]
    for d, sub in f.groupby("design", sort=False):
        s = sub.set_index("rho")
        rows.append([d, str(int(sub.n_groups.iloc[0]))]
                    + [f"{s.loc[x, 'expected_gap']:.3f}"
                       for x in (0.60, 0.70, 0.80)])
    table(doc, rows, [1.48, 0.26, 0.55, 0.55, 0.56])


def tbl_annotation(doc):
    meb = pd.read_csv(FX / "table_label_bias_v2.csv")
    sc = pd.read_csv(SR / "scut_group_decomposition.csv")
    sc = sc[sc.grouping == "subgroup"]
    gc = pd.read_csv(Q / "table_group_ceiling.csv").set_index("ethnicity")
    rows = [["Group", "n", "Raters", "Reliability ρ", "Ceiling √ρ",
             "Spread u"]]
    rows.append(["MEBeauty", "", "", "", "", ""])
    for _, r in meb.sort_values("group").iterrows():
        k = gc.loc[r["group"]]
        rows.append(["  " + MEB_LBL[r["group"]], f"{int(r['n'])}",
                     f"{k['raters_mean']:.0f}", f"{r['ceiling'] ** 2:.3f}",
                     f"{r['ceiling']:.3f}", f"{r['sd_ratio']:.2f}"])
    rows.append(["SCUT-FBP5500", "", "", "", "", ""])
    for _, r in sc.sort_values("group").iterrows():
        rows.append(["  " + SCUT_LBL[r["group"]], f"{int(r['n']):,}", "60",
                     f"{r['reliability']:.3f}", f"{r['ceiling']:.3f}",
                     f"{r['u']:.2f}"])
    table(doc, rows, [1.02, 0.48, 0.42, 0.58, 0.52, 0.46], size=7.2,
          rules=(1, 2, 9))


def tbl_parity(doc):
    fa = pd.read_csv(Q / "table_final_fairness.csv")
    fa = fa[fa.grouping == "subgroup"].set_index("model")
    rows = [["Predictor", "Prediction gap", "Error gap",
             "ε² pred.", "Sig. pairs"]]
    for m, lab in (("geometry", "Geometry"), ("appearance", "Appearance"),
                   ("fusion", "Fusion")):
        r = fa.loc[m]
        rows.append([lab, f"{r['pred_gap']:.3f}", f"{r['error_gap']:.3f}",
                     f"{r['eps2_pred']:.3f}", r["sig_pred_pairs"]])
    rows.append(["Human labels", f"{fa.loc['fusion', 'label_gap']:.3f}",
                 "—", "—", "—"])
    table(doc, rows, [0.92, 0.70, 0.56, 0.60, 0.62])


TABLES = {
    "validation": (tbl_validation, "Known-answer recovery of the two "
                   "label-side estimators. Each row generates data in which "
                   "the true correlation is 0.80 by construction and requires "
                   "the estimator to return it; n = 200,000 per condition. "
                   "All six recoveries land within 0.002 of the target."),
    "floor": (tbl_floor, "The sampling-noise floor: the expected max−min "
              "spread in Pearson r when every group has identical true skill "
              "and identical annotation quality, at the observed group sizes. "
              "ρ is the common true skill level; 4,000 replicates each. "
              "These are the values an audit of each design will report when "
              "there is nothing to find."),
    "annotation": (tbl_annotation, "Per-group annotation properties. "
                   "Reliability is that of the mean of the group's ratings; "
                   "the ceiling is the highest correlation any predictor can "
                   "attain against such a label; u is the ratio of the pooled "
                   "label standard deviation to the group's. MEBeauty's "
                   "parameters vary widely across groups and "
                   "SCUT-FBP5500's do not, which is why the two datasets "
                   "decompose differently."),
    "parity": (tbl_parity, "Two parity channels over the same twelve "
               "intersectional cells of MEBeauty, for three predictor "
               "families of increasing accuracy. The prediction gap grows "
               "while the error gap falls, so the two channels certify "
               "different models. The final row gives the spread in the human "
               "labels themselves."),
}


# ------------------------------------------------------------------ parse
def collect_citations(text: str):
    order, seen = [], set()
    for m in re.finditer(r"\[([a-z0-9_]+)\]", text):
        k = m.group(1)
        if k not in seen:
            seen.add(k)
            order.append(k)
    return order


def build() -> Path:
    src = (Path(__file__).resolve().parent / "manuscript.txt").read_text()
    lines = [ln.rstrip("\n") for ln in src.split("\n")]

    meta = {}
    body = []
    for ln in lines:
        if ln.startswith(("TITLE ", "AUTHORS ", "ABSTRACT ", "TERMS ")):
            k, v = ln.split(" ", 1)
            meta[k] = v
        elif ln.strip():
            body.append(ln)

    cites = collect_citations("\n".join(body))
    missing = [c for c in cites if c not in REFS]
    if missing:
        raise SystemExit(f"undefined citation keys: {missing}")
    num = {k: i + 1 for i, k in enumerate(cites)}

    def subst(t: str) -> str:
        t = re.sub(r"\[([a-z0-9_]+)\]", lambda m: f"[{num[m.group(1)]}]", t)
        return t.replace("ETHICSPLACEHOLDER", ETHICS)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = SERIF
    st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.0

    s0 = doc.sections[0]
    _margins(s0)
    _cols(s0, 1)

    _para(doc, meta["TITLE"], 20, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    _para(doc, meta["AUTHORS"], 11, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    _para(doc, AFFIL, 8, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

    s1 = doc.add_section(WD_SECTION.CONTINUOUS)
    _margins(s1)
    _cols(s1, 2)

    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    _runs(p, "Abstract—", 9, bold=True, italic=True)
    _runs(p, meta["ABSTRACT"], 9, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(8)
    _runs(p, "Index Terms—", 9, bold=True, italic=True)
    _runs(p, meta["TERMS"], 9, bold=True)

    sec_i = 0
    sub_i = 0
    roman = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
    letters = "ABCDEFGHIJ"
    tbl_i = 0
    first_body = True

    for ln in body:
        if ln.startswith("#1 "):
            sec_i += 1
            sub_i = 0
            _para(doc, f"{roman[sec_i - 1]}.  {ln[3:]}", 10,
                  align=WD_ALIGN_PARAGRAPH.CENTER, before=10, after=4,
                  caps=True, keep=True)
            first_body = True
        elif ln.startswith("#2 "):
            _para(doc, f"{letters[sub_i]}. {ln[3:]}", 10,
                  align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=2,
                  italic=True, keep=True)
            sub_i += 1
            first_body = True
        elif ln.startswith("p "):
            _para(doc, subst(ln[2:]), 10, indent=0.0 if first_body else 0.2,
                  after=4)
            first_body = False
        elif ln.startswith("@fig "):
            name, w, lab, cap = ln[5:].split("|", 3)
            figure(doc, name, subst(cap), lab, width=float(w))
            first_body = True
        elif ln.startswith("@wfig "):
            name, w, lab, cap = ln[6:].split("|", 3)
            wide_block(doc, lambda d, n=name, w=w, l=lab, c=cap:
                       figure(d, n, subst(c), l, width=float(w)))
            first_body = True
        elif ln.startswith("@eq "):
            body_eq, numb = ln[4:].rsplit("|", 1)
            equation(doc, body_eq, numb.strip())
            first_body = False
        elif ln.startswith("@tbl "):
            key = ln[5:].strip()
            fn, cap = TABLES[key]
            tbl_i += 1
            _para(doc, f"TABLE {roman[tbl_i - 1]}", 8,
                  align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=1,
                  keep=True)
            _para(doc, cap, 8, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=3,
                  keep=True)
            fn(doc)
            _para(doc, "", 6, after=6)
            first_body = True

    # --------------------------------------------------------- references
    _para(doc, "References", 10, align=WD_ALIGN_PARAGRAPH.CENTER, before=10,
          after=4, caps=True, keep=True)
    for k in cites:
        pr = doc.add_paragraph()
        pr.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pr.paragraph_format.left_indent = Inches(0.24)
        pr.paragraph_format.first_line_indent = Inches(-0.24)
        pr.paragraph_format.tab_stops.add_tab_stop(Inches(0.24))
        pr.paragraph_format.space_after = Pt(1)
        pr.paragraph_format.line_spacing = 1.0
        _runs(pr, f"[{num[k]}]\t{REFS[k]}", 8)

    # ---------------------------------------------------------- biographies
    _para(doc, "", 8, after=6)
    for name, bio in BIOS:
        pb = doc.add_paragraph()
        pb.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pb.paragraph_format.space_after = Pt(4)
        pb.paragraph_format.line_spacing = 1.0
        _runs(pb, name + " ", 8, caps=True, bold=True)
        _runs(pb, bio, 8)

    path = OUT / "Manuscript_TTS.docx"
    doc.save(str(path))
    words = sum(len(ln.split()) for ln in body if ln.startswith("p "))
    print(f"  {len(cites)} references, {tbl_i} tables, ~{words} body words")
    return path


def to_pdf(path: Path) -> None:
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir",
         str(path.parent), str(path)],
        check=True, capture_output=True, timeout=300)


if __name__ == "__main__":
    p = build()
    to_pdf(p)
    pdf = p.with_suffix(".pdf")
    n = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True)
    print([l for l in n.stdout.split("\n") if l.startswith("Pages")])
